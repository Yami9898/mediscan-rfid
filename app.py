from flask import Flask, request, jsonify, render_template, session, send_from_directory
import mysql.connector
from flask_cors import CORS
import datetime
from flask import request
from docx import Document
import os
import uuid
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import serial
import time

# --- CONFIG SERIAL POUR ARDUINO ---
# Vérifie d'abord ton port avec : python -m serial.tools.list_ports
# Sur Mac ça ressemble souvent à : /dev/tty.usbmodemXXXX ou /dev/tty.usbserialXXXX
SERIAL_PORT = os.environ.get("SERIAL_PORT", "/dev/cu.usbmodem11401")  # ⚠️ à adapter selon ton Mac
BAUD_RATE = 9600

try:
    ser = serial.Serial(SERIAL_PORT, BAUD_RATE, timeout=1)
    time.sleep(2)  # attendre le reset Arduino
except Exception as e:
    print("⚠️ Impossible d'ouvrir le port série :", e)
    ser = None


app = Flask(__name__)
CORS(app)  # Autorise les appels JS
app.secret_key = os.environ.get("SECRET_KEY", "un_secret_ultra_complexe")  # Nécessaire pour la session

def get_connection():
    return mysql.connector.connect(
        host=os.environ.get("DB_HOST", "localhost"),
        port=int(os.environ.get("DB_PORT", 8889)),
        user=os.environ.get("DB_USER", "root"),
        password=os.environ.get("DB_PASSWORD", "root"),  # Mot de passe MAMP
        database=os.environ.get("DB_NAME", "ehealth")
    )



@app.route("/")
def patient():
    return render_template("patient.html",patient=session.get("patient"))

# Page RDV
@app.route("/rdv")
def rdv():
    # On récupère juste le patient en session
    patient = session.get("patient")
    # La page affichera un tableau vide au départ et JS ira chercher les rdvs via /api/rdv
    return render_template("rdv.html", patient=patient)


@app.route('/dossier')
def dossier_med():
    # On récupère juste le patient en session
    patient = session.get("patient")
    # La page affichera un tableau vide au départ et JS ira chercher les consultations via /api/consultation
    return render_template("dossier_med.html", patient=patient)

@app.route('/examens')
def examens():
    return render_template('examens.html', patient=session.get("patient"))

@app.route('/suivi_medical')
def suivi_medical():
    return render_template('suivi_medical.html', patient=session.get("patient"))

@app.route('/vaccins')
def vaccins():
    return render_template('vaccins.html', patient=session.get("patient"))

@app.route("/get_patient", methods=["POST"])
def get_patient():
    data = request.get_json()
    print("Données reçues :", data)  # <-- Ajoute ça
    patient_id = data.get("id")

    try:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT id, civilite, nom, prenom,date_naissance,sexe, nationalite, ville, adresse,
                    telephone, e_mail, num_cin, prise_en_charge, photo_url
            FROM patients
            WHERE id = %s
        """, (patient_id,))
        row = cursor.fetchone()
        cursor.close()
        conn.close()

        if row:
            patient = {
                "id": row["id"],
                "civilite": row["civilite"],
                "nom": row["nom"],
                "prenom": row["prenom"],
                "date_naissance": row["date_naissance"].strftime("%d/%m/%Y"),
                "sexe": row["sexe"],
                "nationalite": row["nationalite"],
                "ville": row["ville"],
                "adresse": row["adresse"],
                "telephone": row["telephone"],
                "email": row["e_mail"],
                "num_cin": row["num_cin"],
                "prise_en_charge": row["prise_en_charge"],
                "photo_url": row["photo_url"] if row["photo_url"] else "static/photos/default.jpg"
            }
            # On stocke le patient dans la session
            session["patient"] = patient
            return jsonify(patient)
        else:
            return jsonify({"error": "Patient non trouvé."})

    except Exception as e:
        return jsonify({"error": str(e)})
    

@app.route("/scan_card", methods=["POST"])
def scan_card():
    if ser is None:
        return jsonify({"error": "Port série non disponible"})

    try:
        # On demande à l'Arduino de scanner
        ser.write(b"SCAN\n")
        time.sleep(0.5)

        # Lire ce que renvoie l'Arduino
        if ser.in_waiting > 0:
            card_id = ser.readline().decode().strip()
            print("Carte lue :", card_id)

            # Chercher le patient en base avec l’UID de la carte
            conn = get_connection()
            cursor = conn.cursor(dictionary=True)
            cursor.execute("""
                SELECT id, civilite, nom, prenom, date_naissance, sexe, nationalite, ville, adresse,
                       telephone, e_mail, num_cin, prise_en_charge, photo_url
                FROM patients
                WHERE id = %s
            """, (card_id,))
            row = cursor.fetchone()
            cursor.close()
            conn.close()

            if row:
                patient = {
                    "id": row["id"],
                    "civilite": row["civilite"],
                    "nom": row["nom"],
                    "prenom": row["prenom"],
                    "date_naissance": row["date_naissance"].strftime("%d/%m/%Y"),
                    "sexe": row["sexe"],
                    "nationalite": row["nationalite"],
                    "ville": row["ville"],
                    "adresse": row["adresse"],
                    "telephone": row["telephone"],
                    "email": row["e_mail"],
                    "num_cin": row["num_cin"],
                    "prise_en_charge": row["prise_en_charge"],
                    "photo_url": row["photo_url"] if row["photo_url"] else "static/photos/default.jpg"
                }
                session["patient"] = patient
                return jsonify({"status": "ok", "patient": patient})
            else:
                return jsonify({"error": "Aucun patient trouvé pour cette carte."})

        else:
            return jsonify({"error": "Nouvelle carte non détectée."})

    except Exception as e:
        return jsonify({"error": str(e)})


@app.route("/get_patient_session")
def get_patient_session():
    patient = session.get("patient")
    if patient:
        return jsonify(patient)
    else:
        return jsonify({"error": "Aucun patient en session."})
    
import traceback

@app.route("/api/rdv")
def api_rdv():
    patient = session.get("patient")
    if not patient:
        return jsonify([])  # Pas de patient → liste vide
    try:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT num_rdv, date, heure, lieu, motif, type, statut
            FROM rdv
            WHERE id_patient = %s
            ORDER BY date , heure 
        """, (patient["id"],))
        rdvs = cursor.fetchall()
        cursor.close()
        conn.close()

        for r in rdvs:
            # Conversion date
            if isinstance(r["date"], (datetime.date, datetime.datetime)):
                r["date"] = r["date"].strftime("%d/%m/%Y")
            # Conversion heure (time ou timedelta)
            if isinstance(r["heure"], datetime.time):
                r["heure"] = r["heure"].strftime("%H:%M")
            elif isinstance(r["heure"], datetime.timedelta):
                total_seconds = int(r["heure"].total_seconds())
                hours = total_seconds // 3600
                minutes = (total_seconds % 3600) // 60
                r["heure"] = f"{hours:02d}:{minutes:02d}"

        return jsonify(rdvs)

    except Exception as e:
        print("Erreur dans /api/rdv :", e)
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500
    

# Ajouter un rdv
@app.route("/api/rdv", methods=["POST"])
def ajouter_rdv():
    patient = session.get("patient")
    if not patient:
        return jsonify({"error": "Patient non connecté"}), 401
    data = request.get_json()
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO rdv (date, heure, lieu, motif, type, statut, id_patient)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
        """, (
            data.get("date"),    # Format attendu 'YYYY-MM-DD'
            data.get("heure"),   # Format 'HH:MM:SS'
            data.get("lieu"),
            data.get("motif"),
            data.get("type"),
            data.get("statut"),
            patient["id"]
        ))
        conn.commit()
        cursor.close()
        conn.close()
        return jsonify({"message": "Rendez-vous ajouté avec succès"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    
@app.route("/api/rdv/<int:num_rdv>", methods=["PUT"])
def modifier_rdv(num_rdv):
    patient = session.get("patient")
    if not patient:
        return jsonify({"error": "Patient non connecté"}), 401
    data = request.get_json()
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            UPDATE rdv
            SET date=%s, heure=%s, lieu=%s, motif=%s, type=%s, statut=%s
            WHERE num_rdv=%s AND id_patient=%s
        """, (
            data.get("date"),
            data.get("heure"),
            data.get("lieu"),
            data.get("motif"),
            data.get("type"),
            data.get("statut"),
            num_rdv,
            patient["id"]
        ))
        conn.commit()
        cursor.close()
        conn.close()
        return jsonify({"message": "Rendez-vous modifié avec succès"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# Supprimer un rdv
@app.route("/api/rdv/<int:num_rdv>", methods=["DELETE"])
def supprimer_rdv(num_rdv):
    patient = session.get("patient")
    if not patient:
        return jsonify({"error": "Patient non connecté"}), 401
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            DELETE FROM rdv
            WHERE num_rdv=%s AND id_patient=%s
        """, (num_rdv, patient["id"]))
        conn.commit()
        cursor.close()
        conn.close()
        return jsonify({"message": "Rendez-vous supprimé avec succès"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    
@app.route('/api/rdv/search')
def search_rdv():
    # prend q comme paramètre de recherche
    q = request.args.get('q', '').strip()
    patient = session.get("patient")
    if not patient:
        return jsonify([])  # pas de patient connecté -> vide

    try:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)

        # On recherche seulement les rdv du patient
        sql = """
            SELECT num_rdv, date, heure, lieu, type, type_imagerie, statut
            FROM rdv
            WHERE id_patient = %s
              AND (
                    CAST(date AS CHAR) LIKE %s
                 OR CAST(heure AS CHAR) LIKE %s
                 OR lieu LIKE %s
                 OR type LIKE %s
                 OR type_imagerie LIKE %s
                 OR statut LIKE %s
              )
            ORDER BY date DESC, heure DESC
        """
        like = f"%{q}%"
        params = (patient["id"], like, like, like, like, like, like)
        cursor.execute(sql, params)
        rdvs = cursor.fetchall()

        # conversion cohérente date/heure comme dans /api/rdv
        import datetime
        for r in rdvs:
            if isinstance(r.get("date"), (datetime.date, datetime.datetime)):
                r["date"] = r["date"].strftime("%d/%m/%Y")
            if isinstance(r.get("heure"), datetime.time):
                r["heure"] = r["heure"].strftime("%H:%M")
            elif isinstance(r.get("heure"), datetime.timedelta):
                total_seconds = int(r["heure"].total_seconds())
                hours = total_seconds // 3600
                minutes = (total_seconds % 3600) // 60
                r["heure"] = f"{hours:02d}:{minutes:02d}"

        cursor.close()
        conn.close()
        return jsonify(rdvs)
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    

ORDONNANCES_FOLDER = "static/ordonnances"
def generer_ordonnance(consultation, patient):
    """Génère un fichier docx d'ordonnance et retourne le chemin"""
    if not os.path.exists(ORDONNANCES_FOLDER):
        os.makedirs(ORDONNANCES_FOLDER)
    
    filename = f"ordonnance_{consultation['num_consultation']}_{uuid.uuid4().hex[:6]}.docx"
    filepath = os.path.join(ORDONNANCES_FOLDER, filename)
    
    doc = Document()
    
    # Configuration de base
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)
    
    ## En-tête professionnel
    header = doc.sections[0].header
    header_para = header.paragraphs[0]
    header_para.text = f"{consultation['etablissement']} • {consultation['medecin']} • {consultation['specialite']}"
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    ## Titre principal
    title = doc.add_heading('ORDONNANCE MÉDICALE', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.style.font.color.rgb = RGBColor(0, 0, 0)
    
    ## Information patient
    doc.add_paragraph(" ", style='Normal')
    patient_info = doc.add_paragraph()
   
    patient_info.add_run(f"{patient['civilite']} {patient['nom']} {patient['prenom']}")
    
    # Convertir la date avant l'affichage
    from datetime import datetime
    # Format initial supposé : '2023-12-31' (YYYY-MM-DD)
    date_consultation = consultation['date']  # ou consultation.get('date')
    # Conversion en objet datetime si nécessaire
    if isinstance(date_consultation, str):
        try:
            date_obj = datetime.strptime(date_consultation, '%Y-%m-%d')
            date_formatee = date_obj.strftime('%d/%m/%Y')
        except ValueError:
            date_formatee = date_consultation  # Garde la valeur originale si échec
    else:
        date_formatee = date_consultation.strftime('%d/%m/%Y')
    
    patient_info.add_run("\nDate de consultation: ").bold = True
    patient_info.add_run(f"{date_formatee}")
    
    ## Section de prescription
    doc.add_paragraph(" ", style='Normal')
    prescription_title = doc.add_paragraph()
    prescription_title.add_run("PRESCRIPTION MÉDICALE").bold = True
    prescription_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Ajout du contenu traitement depuis la base de données (clé 'traitement' du dictionnaire consultation)
    if consultation.get('traitement'):
        # Convertit les sauts de ligne en paragraphes Word
        for line in consultation['traitement'].split('\n'):
            if line.strip():  # Ne crée pas de paragraphe pour les lignes vides
                doc.add_paragraph(line.strip(), style='Normal')
    else:
        doc.add_paragraph("Aucun traitement prescrit", style='Normal')
    
    ## Signature
    doc.add_paragraph(" ", style='Normal')
    doc.add_paragraph(" ", style='Normal')
    signature_line = doc.add_paragraph()
    signature_line.add_run("_" * 30).bold = True
    signature_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    signature_text = doc.add_paragraph()
    signature_text.add_run(consultation['medecin']).bold = True
    signature_text.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    # Sauvegarde
    filename = f"ordonnance_{consultation['num_consultation']}.docx"
    filepath = os.path.join("static","ordonnances", filename)
    doc.save(filepath)
    
    return f"static/ordonnances/{filename}"

    
    doc.save(filepath)
    return filepath.replace("static/", "")



@app.route('/api/generer-ordonnance/<int:consult_id>')
def generer_ordonnance_route(consult_id):
    try:
        # 1. Récupérer la consultation ET le patient en une seule requête
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)
        
        cursor.execute("""
            SELECT c.*, p.* 
            FROM consultations c
            JOIN patients p ON c.id_patient = p.id
            WHERE c.num_consultation = %s
        """, (consult_id,))
        
        data = cursor.fetchone()
        
        if not data:
            return jsonify({"error": "Consultation non trouvée"}), 404

        # 2. Générer l'ordonnance
        filepath = generer_ordonnance(data, data)  # data contient consultation + patient
        
        # 3. Renvoyer le fichier
        directory = os.path.dirname(filepath)
        filename = os.path.basename(filepath)
        
        return send_from_directory(directory, filename, as_attachment=True)
        
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        cursor.close()
        conn.close()

    
   

@app.route("/api/consultation", methods=["GET"])
def api_consultation():
    patient = session.get("patient")
    if not patient:
        return jsonify([]) # Pas de patient → liste vide 
    try:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT num_consultation,date,etablissement, medecin, specialite, maladie, traitement, statut, pv
            FROM consultations
            WHERE id_patient = %s
            ORDER BY date DESC
        """, (patient["id"],))
        consultations = cursor.fetchall()
        cursor.close()
        conn.close()

        for r in consultations:
            # Conversion date
            if isinstance(r["date"], (datetime.date, datetime.datetime)):
                r["date"] = r["date"].strftime("%d/%m/%Y")

        return jsonify(consultations)

    except Exception as e:
        print("Erreur dans /api/consultation :", e)
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500
    
@app.route("/api/consultation", methods=["POST"])
# Ajouter une consultation
def ajouter_consultation():
    patient = session.get("patient")
    if not patient:
        return jsonify({"error": "Patient non connecté"}), 401
    data = request.get_json()
    try:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            INSERT INTO consultations (date, etablissement, medecin, specialite, maladie, traitement, statut, pv, id_patient)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
        """, (
            data.get("date"),    # Format attendu 'YYYY-MM-DD'
            data.get("etablissement"),
            data.get("medecin"),
            data.get("specialite"),
            data.get("maladie"),
            data.get("traitement"),
            data.get("statut"),
            data.get("pv"),
            patient["id"]
        ))

        # Récupérer l'ID de la consultation nouvellement créée
        consult_id = cursor.lastrowid
        
        # Récupérer les infos pour générer l'ordonnance
        cursor.execute("SELECT * FROM consultations WHERE num_consultation = %s", (consult_id,))
        new_consult = cursor.fetchone()
        
        # Générer l'ordonnance
        ordonnance_path = generer_ordonnance(new_consult, patient)
        
        # Mettre à jour le champ ordonnance
        cursor.execute("""
            UPDATE consultations 
            SET ordonnance = %s 
            WHERE num_consultation = %s
        """, (ordonnance_path, consult_id))


        conn.commit()
        cursor.close()
        conn.close()
        return jsonify({"message": "Consultation ajoutée avec succès"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    
@app.route("/api/consultation/<int:num_consultation>", methods=["PUT"])
def modifier_consultation(num_consultation):
    patient = session.get("patient")
    if not patient:
        return jsonify({"error": "Patient non connecté"}), 401
    data = request.get_json()




    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            UPDATE consultations
            SET date=%s, etablissement=%s, medecin=%s, specialite=%s, maladie=%s, traitement=%s, statut=%s, pv=%s
            WHERE num_consultation=%s AND id_patient=%s
        """, (
            data.get("date"),
            data.get("etablissement"),
            data.get("medecin"),
            data.get("specialite"),
            data.get("maladie"),
            data.get("traitement"),
            data.get("statut"),
            data.get("pv"),
            num_consultation,
            patient["id"]
        ))
        conn.commit()
        cursor.close()
        conn.close()
        return jsonify({"message": "Consultation modifiée avec succès"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    
    
# Supprimer une consultation
@app.route("/api/consultation/<int:num_consultation>", methods=["DELETE"])
def supprimer_consultation(num_consultation):
    patient = session.get("patient")
    if not patient:
        return jsonify({"error": "Patient non connecté"}), 401
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            DELETE FROM consultations
            WHERE num_consultation=%s AND id_patient=%s
        """, (num_consultation, patient["id"]))
        conn.commit()
        cursor.close()
        conn.close()
        return jsonify({"message": "Consultation supprimée avec succès"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    
@app.route('/api/consultation/search')
def search_consultation():
    # prend q comme paramètre de recherche
    q = request.args.get('q', '').strip()
    patient = session.get("patient")
    if not patient:
        return jsonify([])  # pas de patient connecté -> vide

    try:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)

        # On recherche seulement les consultations du patient
        sql = """
            SELECT num_consultation, date, heure, medecin, specialite, ordonnance, pv
            FROM consultations
            WHERE id_patient = %s
              AND (
                    CAST(date AS CHAR) LIKE %s
                 OR etablissement LIKE %s
                 OR medecin LIKE %s
                 OR specialite LIKE %s
                 OR maladie LIKE %s
                 OR traitement LIKE %s
                 OR statut LIKE %s
                 OR pv LIKE %s
              )
            ORDER BY date DESC
        """
        like = f"%{q}%"
        params = (patient["id"], like, like, like, like, like, like, like, like)
        cursor.execute(sql, params)
        consultations = cursor.fetchall()

        # conversion cohérente date/heure comme dans /api/cosultation
        import datetime
        for r in consultations:
            if isinstance(r.get("date"), (datetime.date, datetime.datetime)):
                r["date"] = r["date"].strftime("%d/%m/%Y")
           

        cursor.close()
        conn.close()
        return jsonify(consultations)
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    

@app.route("/api/vaccins")
def api_vaccins():
    patient = session.get("patient")
    if not patient:
        return jsonify([])  # Pas de patient → liste vide
    try:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT num_vaccin, nom, type, date, dose, voie, lieu, responsable
            FROM vaccins
            WHERE id_patient = %s
            ORDER BY date DESC
        """, (patient["id"],))
        vaccins = cursor.fetchall()
        cursor.close()
        conn.close()

        for r in vaccins:
            # Conversion date
            if isinstance(r["date"], (datetime.date, datetime.datetime)):
                r["date"] = r["date"].strftime("%d/%m/%Y")
            
        return jsonify(vaccins)

    except Exception as e:
        print("Erreur dans /api/vaccins :", e)
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

 # Ajouter un vaccin
@app.route("/api/vaccins", methods=["POST"])
def ajouter_vaccin():
    patient = session.get("patient")
    if not patient:
        return jsonify({"error": "Patient non connecté"}), 401
    data = request.get_json()
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO vaccins (nom, type, date, dose, voie, lieu, responsable, id_patient)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
        """, (
            data.get("nom"),    
            data.get("type"),   
            data.get("date"), # Format attendu 'YYYY-MM-DD'
            data.get("dose"),
            data.get("voie"),
            data.get("lieu"),
            data.get("responsable"),
            patient["id"]
        ))
        conn.commit()
        cursor.close()
        conn.close()
        return jsonify({"message": "Vaccin ajouté avec succès"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    
@app.route("/api/vaccins/<int:num_vaccin>", methods=["PUT"])
def modifier_vaccin(num_vaccin):
    patient = session.get("patient")
    if not patient:
        return jsonify({"error": "Patient non connecté"}), 401
    data = request.get_json()
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            UPDATE vaccins
            SET nom=%s, type=%s, date=%s, dose=%s, voie=%s, lieu=%s , responsable=%s
            WHERE num_vaccin=%s AND id_patient=%s
        """, (
            data.get("nom"),
            data.get("type"),
            data.get("date"),
            data.get("dose"),
            data.get("voie"),
            data.get("lieu"),
            data.get("responsable"),
            num_vaccin,
            patient["id"]
        ))
        conn.commit()
        cursor.close()
        conn.close()
        return jsonify({"message": "Vaccin modifié avec succès"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    
# Supprimer un vaccin
@app.route("/api/vaccins/<int:num_vaccin>", methods=["DELETE"])
def supprimer_vaccin(num_vaccin):
    patient = session.get("patient")
    if not patient:
        return jsonify({"error": "Patient non connecté"}), 401
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            DELETE FROM vaccins
            WHERE num_vaccin=%s AND id_patient=%s
        """, (num_vaccin, patient["id"]))
        conn.commit()
        cursor.close()
        conn.close()
        return jsonify({"message": "Vaccin supprimé avec succès"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    
@app.route('/api/vaccins/search')
def search_vaccin():
    # prend q comme paramètre de recherche
    q = request.args.get('q', '').strip()
    patient = session.get("patient")
    if not patient:
        return jsonify([])  # pas de patient connecté -> vide

    try:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)

        # On recherche seulement les vaccins du patient
        sql = """
            SELECT num_vaccin, nom, type, date, dose, voie, lieu, responsable
            FROM vaccins
            WHERE id_patient = %s
              AND (
                    CAST(date AS CHAR) LIKE %s
                 OR nom LIKE %s
                 OR type LIKE %s
                 OR dose LIKE %s
                 OR voie LIKE %s
                 OR lieu LIKE %s
                 OR responsable LIKE %s
              )
            ORDER BY date DESC
        """
        like = f"%{q}%"
        params = (patient["id"], like, like, like, like, like, like, like)
        cursor.execute(sql, params)
        vaccins = cursor.fetchall()

        # conversion cohérente date comme dans /api/vaccins
        import datetime
        for r in vaccins:
            if isinstance(r.get("date"), (datetime.date, datetime.datetime)):
                r["date"] = r["date"].strftime("%d/%m/%Y")
            

        cursor.close()
        conn.close()
        return jsonify(vaccins)
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    

#page de suivi médical

@app.route("/api/consultations_en_cours")
def api_consultations_en_cours():
    patient = session.get("patient")
    if not patient:
        return jsonify([])  # Pas de patient → liste vide
    try:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT num_consultation, maladie, date, traitement, pv, statut
            FROM consultations
            WHERE id_patient = %s AND statut = 'en cours'
        """, (patient["id"],))
        consultations = cursor.fetchall()
        cursor.close()
        conn.close()
        for c in consultations:
            # Conversion de la date si nécessaire
            if isinstance(c["date"], (datetime.date, datetime.datetime)):
                c["date"] = c["date"].strftime("%d/%m/%Y")
        return jsonify(consultations)
    
    except Exception as e:
        print("Erreur dans /api/consultations_en_cours :", e)
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500
    

    
@app.route("/api/update_statut", methods=["POST"]) 
def update_statut():
    patient = session.get("patient")
    if not patient:
        return jsonify([])  # Pas de patient → liste vide
    try:
        data = request.get_json()
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            UPDATE consultations
            SET statut = %s
            WHERE num_consultation = %s
        """, (data['new_statut'], data['id']))  # Assurez-vous que 'id' correspond à num_consultation
        conn.commit()
        cursor.close()
        conn.close()
        return jsonify({"success": True})
    except Exception as e:
        print("Erreur lors de la mise à jour:", e)
        return jsonify({"error": str(e)}), 500


@app.route("/api/maladies_chroniques")
def api_maladies_chroniques():
    patient = session.get("patient")
    if not patient:
        return jsonify([])  # Pas de patient → liste vide
    try:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT num_maladie, maladie, type, historique_trait, avancement, date_debut, date_fin
            FROM maladies
            WHERE id_patient = %s
        """, (patient["id"],))
        maladies = cursor.fetchall()
        cursor.close()
        conn.close()

        for m in maladies:
            # Conversion des dates si nécessaire
            if isinstance(m["date_debut"], (datetime.date, datetime.datetime)):
                m["date_debut"] = m["date_debut"].strftime("%d/%m/%Y")
            if isinstance(m["date_fin"], (datetime.date, datetime.datetime)):
                m["date_fin"] = m["date_fin"].strftime("%d/%m/%Y")

        return jsonify(maladies)

    except Exception as e:
        print("Erreur dans /api/maladies_chroniques :", e)
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

# Ajouter une maladie chronique
@app.route("/api/maladies", methods=["POST"])
def ajouter_maladie():
    patient = session.get("patient")
    if not patient:
        return jsonify({"error": "Patient non connecté"}), 401
    data = request.get_json()
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO maladies (maladie, type, historique_trait, avancement, date_debut, date_fin, id_patient)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
        """, (
            data.get("maladie"),
            data.get("type"),
            data.get("historique_trait"),
            data.get("avancement"),
            data.get("date_debut"),
            data.get("date_fin"),
            patient["id"]  # Utiliser l'ID du patient de la session
        ))
        conn.commit()
        cursor.close()
        conn.close()
        return jsonify({"message": "Maladie ajoutée avec succès"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

# Modifier une maladie chronique
@app.route("/api/maladies/<int:num_maladie>", methods=["PUT"])
def modifier_maladie(num_maladie):
    patient = session.get("patient")
    if not patient:
        return jsonify({"error": "Patient non connecté"}), 401
    data = request.get_json()
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            UPDATE maladies
            SET maladie=%s, type=%s, historique_trait=%s, avancement=%s, date_debut=%s, date_fin=%s
            WHERE num_maladie=%s AND id_patient=%s
        """, (
            data.get("maladie"),
            data.get("type"),
            data.get("historique_trait"),
            data.get("avancement"),
            data.get("date_debut"),
            data.get("date_fin"),
            num_maladie,
            patient["id"]
        ))
        conn.commit()
        cursor.close()
        conn.close()
        return jsonify({"message": "Maladie modifiée avec succès"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

# Supprimer une maladie chronique
@app.route("/api/maladies/<int:num_maladie>", methods=["DELETE"])
def supprimer_maladie(num_maladie):
    patient = session.get("patient")
    if not patient:
        return jsonify({"error": "Patient non connecté"}), 401
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            DELETE FROM maladies
            WHERE num_maladie=%s AND id_patient=%s
        """, (num_maladie, patient["id"]))
        conn.commit()
        cursor.close()
        conn.close()
        return jsonify({"message": "Maladie supprimée avec succès"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500




if __name__ == "__main__":
    app.run(debug=True)

